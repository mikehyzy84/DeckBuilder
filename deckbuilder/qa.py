"""Automated QA for a built deck.

Checks every slide for: margins under 0.5 inch, likely text overflow, overlapping text boxes,
leftover placeholder text, empty placeholders, em or en dashes, missing speaker notes,
accent lines under titles, decorative edge stripes, off-brand colors and fonts, low text contrast,
and pictures that look like charts. Then renders every slide so Claude can inspect the images.

Automated checks catch the mechanical problems. Claude still looks at every rendered slide.
"""
from __future__ import annotations

import json
import re
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_AUTO_SIZE

from . import config as C
from .classify import PLACEHOLDER_RE

EMU = 914400
MIN_MARGIN_IN = 0.5
DASHES = re.compile("[—–]")
EXEMPT_PH = {PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.DATE}


def _lum(hexstr: str) -> float:
    r, g, b = (int(hexstr[i:i + 2], 16) / 255 for i in (0, 2, 4))
    f = lambda c: c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4
    return 0.2126 * f(r) + 0.7152 * f(g) + 0.0722 * f(b)


def contrast(a: str, b: str) -> float:
    la, lb = sorted((_lum(a), _lum(b)), reverse=True)
    return (la + 0.05) / (lb + 0.05)


def _brand(slug: str | None) -> dict:
    if not slug:
        return {}
    p = C.REPO_ROOT / "docs" / "brands" / slug / "brand.json"
    return json.loads(p.read_text()) if p.exists() else {}


def _allowed_colors(brand: dict) -> set[str]:
    cols = set()
    for v in (brand.get("theme", {}).get("colors", {}) or {}).values():
        if v:
            cols.add(v.lstrip("#").upper())
    for v in brand.get("palette_allowed", []):
        cols.add(v.lstrip("#").upper())
    return cols


def _is_exempt(sh) -> bool:
    try:
        return sh.is_placeholder and sh.placeholder_format.type in EXEMPT_PH
    except Exception:
        return False


def _overflow(sh) -> bool:
    tf = sh.text_frame
    if tf.auto_size in (MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT, MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE):
        return False
    if not sh.width or not sh.height:
        return False
    width_pt = sh.width / 12700 - 14
    height_pt = sh.height / 12700 - 7
    need = 0.0
    for p in tf.paragraphs:
        size = next((r.font.size.pt for r in p.runs if r.font.size), 14)
        text = "".join(r.text for r in p.runs)
        cpl = max(1, int(width_pt / (size * 0.5)))
        lines = max(1, -(-len(text) // cpl)) if text else 1
        need += lines * size * 1.2
    return need > height_pt * 1.08


def _boxes_overlap(a, b) -> float:
    ax2, ay2 = a.left + a.width, a.top + a.height
    bx2, by2 = b.left + b.width, b.top + b.height
    w = min(ax2, bx2) - max(a.left, b.left)
    h = min(ay2, by2) - max(a.top, b.top)
    if w <= 0 or h <= 0:
        return 0.0
    return (w * h) / max(1, min(a.width * a.height, b.width * b.height))


def check_deck(deck: Path, brand_slug: str | None = None) -> list[dict]:
    prs = Presentation(str(deck))
    sw, shh = prs.slide_width, prs.slide_height
    brand = _brand(brand_slug)
    allowed = _allowed_colors(brand)
    brand_fonts = {f for f in (brand.get("theme", {}).get("fonts", {}) or {}).values() if f}
    issues: list[dict] = []

    def add(n, kind, detail, shape=None):
        issues.append({"slide": n, "check": kind, "detail": detail, "shape": getattr(shape, "name", None)})

    for n, slide in enumerate(prs.slides, 1):
        texts = []
        title_shape = None
        for sh in slide.shapes:
            try:
                if sh.is_placeholder and sh.placeholder_format.type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
                    title_shape = sh
            except Exception:
                pass
        for sh in slide.shapes:
            if sh.left is None or sh.width is None or _is_exempt(sh):
                continue
            full_bleed = sh.width >= sw * 0.95 or sh.height >= shh * 0.95
            is_pic = sh.shape_type == MSO_SHAPE_TYPE.PICTURE
            has_text = getattr(sh, "has_text_frame", False) and sh.has_text_frame and sh.text_frame.text.strip()
            # margins apply to content, not full-bleed backgrounds or images placed to the edge on purpose
            if has_text and not full_bleed:
                l, t = sh.left / EMU, sh.top / EMU
                r, b = (sw - sh.left - sh.width) / EMU, (shh - sh.top - sh.height) / EMU
                if min(l, t, r, b) < MIN_MARGIN_IN - 0.01:
                    add(n, "margin", f"text box {min(l, t, r, b):.2f} in from edge", sh)
            # accent line under title or decorative stripe
            if sh.shape_type in (MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.LINE) and not has_text:
                thin = min(sh.width, sh.height) < EMU * 0.08 and max(sh.width, sh.height) > EMU * 0.8
                if thin and title_shape is not None and title_shape.top is not None:
                    below = sh.top - (title_shape.top + title_shape.height)
                    if -EMU * 0.05 < below < EMU * 0.4 and sh.width > sh.height:
                        add(n, "accent-line", "thin rule directly under the title", sh)
                if thin and (sh.left <= EMU * 0.05 or sh.top <= EMU * 0.05 or
                             sw - sh.left - sh.width <= EMU * 0.05 or shh - sh.top - sh.height <= EMU * 0.05):
                    add(n, "edge-stripe", "thin decorative shape on the slide edge", sh)
            if is_pic and re.search(r"chart|graph", sh.name, re.I):
                add(n, "chart-image", "picture named like a chart; use a native chart", sh)
            if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                txt = sh.text_frame.text
                if sh.is_placeholder and not txt.strip():
                    add(n, "empty-placeholder", "placeholder left empty; delete it or fill it", sh)
                if not txt.strip():
                    continue
                texts.append((sh, txt))
                if PLACEHOLDER_RE.search(txt):
                    add(n, "placeholder-text", txt[:80], sh)
                if DASHES.search(txt):
                    add(n, "dash", "em or en dash in slide text", sh)
                if _overflow(sh):
                    add(n, "overflow", "text likely exceeds its box", sh)
                fill = None
                try:
                    if sh.fill.type == 1 and sh.fill.fore_color.type == 1:
                        fill = str(sh.fill.fore_color.rgb)
                except Exception:
                    pass
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        if brand_fonts and r.font.name and r.font.name not in brand_fonts and not r.font.name.startswith("+"):
                            add(n, "font", f"font {r.font.name} not in brand fonts {sorted(brand_fonts)}", sh)
                        try:
                            if r.font.color and r.font.color.type == 1:
                                c = str(r.font.color.rgb)
                                if allowed and c.upper() not in allowed:
                                    add(n, "off-palette", f"text color #{c}", sh)
                                if fill and contrast(c, fill) < 4.5:
                                    add(n, "contrast", f"#{c} on #{fill} ratio {contrast(c, fill):.1f} (< 4.5)", sh)
                        except Exception:
                            pass
            try:
                if allowed and sh.fill.type == 1 and sh.fill.fore_color.type == 1:
                    c = str(sh.fill.fore_color.rgb).upper()
                    if c not in allowed:
                        add(n, "off-palette", f"fill color #{c}", sh)
            except Exception:
                pass
        for i in range(len(texts)):
            for j in range(i + 1, len(texts)):
                a, b = texts[i][0], texts[j][0]
                if None in (a.top, b.top, a.width, b.width):
                    continue
                if _boxes_overlap(a, b) > 0.15:
                    add(n, "overlap", f"'{a.name}' overlaps '{b.name}'", a)
        cls = brand.get("footer_classification")
        if cls:
            for sh in list(slide.shapes) + list(slide.slide_layout.shapes) + list(slide.slide_layout.slide_master.shapes):
                if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                    t = sh.text_frame.text.strip()
                    if t in ("Internal", "Confidential", "Public", "Restricted") and t != cls:
                        add(n, "classification", f"footer shows '{t}', brand requires '{cls}'", sh)
        notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide and slide.notes_slide.notes_text_frame else ""
        if not notes:
            add(n, "notes", "no speaker notes")
        elif DASHES.search(notes):
            add(n, "dash", "em or en dash in speaker notes")
    # de-duplicate repeated identical findings on one shape
    seen, out = set(), []
    for i in issues:
        k = (i["slide"], i["check"], i["shape"], i["detail"])
        if k not in seen:
            seen.add(k)
            out.append(i)
    return out


def run(cfg: dict, deck: Path, brand: str | None, render_dir: str | None) -> int:
    from .render import contact_sheet, render_deck
    issues = check_deck(deck, brand)
    qa_dir = Path(render_dir) if render_dir else deck.parent / "qa"
    qa_dir.mkdir(parents=True, exist_ok=True)
    pngs = render_deck(deck, qa_dir / "slides", cfg, width=1600)
    pngs.sort(key=lambda p: int(p.stem))
    prs = Presentation(str(deck))
    if len(pngs) != len(prs.slides):
        issues.append({"slide": 0, "check": "render", "detail": f"rendered {len(pngs)} of {len(prs.slides)} slides", "shape": None})
    contact_sheet(pngs, qa_dir / "contact_sheet.jpg", cols=4)
    (qa_dir / "qa_report.json").write_text(json.dumps(issues, indent=1))
    by = {}
    for i in issues:
        by.setdefault(i["check"], 0)
        by[i["check"]] += 1
    print(f"{len(prs.slides)} slides, {len(issues)} findings: {by}")
    for i in issues:
        print(f"  slide {i['slide']:>3}  {i['check']:<17} {i['detail']}  [{i['shape']}]")
    print(f"renders: {qa_dir/'slides'}  contact sheet: {qa_dir/'contact_sheet.jpg'}")
    return 0 if not issues else 3
