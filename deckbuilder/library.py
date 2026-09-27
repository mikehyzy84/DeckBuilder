"""Master slide database.

Indexes every slide of every deck in the template library and every registered brand folder
into SLIDE_LIBRARY_ROOT/slides.db (SQLite + FTS5), with a thumbnail per slide and a contact
sheet per deck. Paths are stored relative to their source root so the database works on any
machine where config.local.yaml points at the same folders.

Layout of SLIDE_LIBRARY_ROOT:
  slides.db
  catalog.csv                 flat export of the slides table
  thumbs/<deck_id>/<n>.png    640 px wide per slide
  sheets/<deck_id>/<k>.jpg    contact sheets, 48 slides each
"""
from __future__ import annotations

import csv
import datetime as dt
import json
import re
import sqlite3
from collections import Counter
from pathlib import Path

from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.util import Emu

from . import config as C
from .classify import PLACEHOLDER_RE, classify
from .pptx_utils import emu_to_in, file_md5, open_presentation

DECK_EXT = {".pptx", ".potx"}
SCHEMA = """
CREATE TABLE IF NOT EXISTS decks (
  id TEXT PRIMARY KEY, source TEXT, rel_path TEXT, name TEXT, md5 TEXT,
  n_slides INTEGER, width_in REAL, height_in REAL, n_layouts INTEGER,
  indexed_at TEXT, rendered INTEGER DEFAULT 0
);
CREATE TABLE IF NOT EXISTS slides (
  id TEXT PRIMARY KEY, deck_id TEXT, source TEXT, slide_num INTEGER, layout_name TEXT,
  title TEXT, body_text TEXT, notes TEXT, word_count INTEGER, n_shapes INTEGER,
  n_text_shapes INTEGER, n_pictures INTEGER, picture_area_pct REAL, n_charts INTEGER,
  chart_types TEXT, n_tables INTEGER, table_dims TEXT, n_groups INTEGER, n_smartart INTEGER,
  n_connectors INTEGER, n_autoshapes INTEGER, n_arrow_shapes INTEGER, n_freeforms INTEGER,
  autoshape_types TEXT, n_columns_est INTEGER, has_placeholder_text INTEGER, assertion_title INTEGER,
  title_words INTEGER, min_font_pt REAL, max_font_pt REAL, fonts TEXT, colors TEXT,
  margin_left_in REAL, margin_right_in REAL, margin_top_in REAL, margin_bottom_in REAL,
  category TEXT, subcategory TEXT, tags TEXT, curated_group TEXT, hidden INTEGER, thumb TEXT
);
CREATE TABLE IF NOT EXISTS overrides (slide_id TEXT PRIMARY KEY, category TEXT, subcategory TEXT, tags TEXT, note TEXT, quality INTEGER);
CREATE TABLE IF NOT EXISTS layouts (
  deck_id TEXT, master_idx INTEGER, layout_idx INTEGER, name TEXT, placeholders TEXT,
  PRIMARY KEY (deck_id, master_idx, layout_idx)
);
CREATE VIRTUAL TABLE IF NOT EXISTS slides_fts USING fts5(id UNINDEXED, title, body_text, notes, category, tags, layout_name);
CREATE INDEX IF NOT EXISTS ix_slides_cat ON slides(category);
CREATE INDEX IF NOT EXISTS ix_slides_deck ON slides(deck_id);
"""


def connect(cfg: dict) -> sqlite3.Connection:
    db = sqlite3.connect(C.library_root(cfg) / "slides.db")
    db.row_factory = sqlite3.Row
    db.executescript(SCHEMA)
    return db


def slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")


def deck_id_for(source: str, rel: Path) -> str:
    """Short stable id: sw-<file stem> for the template library, <brand>-<file stem> for brands."""
    prefix = "sw" if source == "slideworks" else source.replace("brand:", "")
    stem = re.sub(r"^slideworks[-_ ]*", "", rel.stem, flags=re.I)
    ext = "-potx" if rel.suffix.lower() == ".potx" else ""
    return slug(f"{prefix}-{stem}{ext}")


def find_decks(cfg: dict) -> list[tuple[str, Path, Path]]:
    import fnmatch
    out = []
    excludes = cfg.get("LIBRARY_EXCLUDE") or []
    for source, root in C.source_roots(cfg).items():
        if not root.exists():
            print(f"WARNING missing source folder {root}")
            continue
        for p in sorted(root.rglob("*")):
            if any(fnmatch.fnmatch(p.name, pat) for pat in excludes):
                continue
            if p.suffix.lower() in DECK_EXT and not p.name.startswith("~$") and p.is_file():
                out.append((source, root, p.relative_to(root)))
    return out


# ---------- feature extraction ----------

def _iter_shapes(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(sh.shapes)


def _text_of(sh) -> str:
    if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
        return sh.text_frame.text.replace("\x0b", " ").strip()
    return ""


def _title(slide) -> str:
    for sh in slide.placeholders:
        try:
            if sh.placeholder_format.type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
                t = _text_of(sh)
                if t:
                    return t
        except Exception:
            continue
    # fallback: the largest-font text near the top
    best, best_key = "", None
    for sh in slide.shapes:
        t = _text_of(sh)
        if not t or sh.top is None:
            continue
        size = 0
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if r.font.size:
                    size = max(size, r.font.size.pt)
        key = (-(size or 0), sh.top)
        if sh.top < Emu(914400 * 1.5) and (best_key is None or key < best_key):
            best, best_key = t, key
    return best


def _columns(text_boxes, slide_w) -> int:
    """Estimate side-by-side text columns: boxes that share a top band and differ in left."""
    bands: dict[int, set[int]] = {}
    for sh in text_boxes:
        if sh.width is None or sh.width > slide_w * 0.6:
            continue
        band = int(emu_to_in(sh.top) * 2)
        bands.setdefault(band, set()).add(int(emu_to_in(sh.left) * 2))
    return max((len(v) for v in bands.values()), default=0)


def extract(slide, n: int, sw: int, shh: int) -> dict:
    counts = Counter()
    ashape_types = Counter()
    chart_types, table_dims, fonts, colors, sizes = [], [], Counter(), Counter(), []
    texts, text_boxes = [], []
    pic_area = 0
    left = top = None
    right = bottom = None
    for sh in _iter_shapes(slide.shapes):
        counts["shapes"] += 1
        st = sh.shape_type
        if sh.left is not None and sh.width and sh.top is not None and sh.height:
            left = sh.left if left is None else min(left, sh.left)
            top = sh.top if top is None else min(top, sh.top)
            r, b = sh.left + sh.width, sh.top + sh.height
            right = r if right is None else max(right, r)
            bottom = b if bottom is None else max(bottom, b)
        if st == MSO_SHAPE_TYPE.GROUP:
            counts["groups"] += 1
            continue
        if st == MSO_SHAPE_TYPE.PICTURE or (st == MSO_SHAPE_TYPE.PLACEHOLDER and sh.__class__.__name__ == "PlaceholderPicture"):
            counts["pictures"] += 1
            if sh.width and sh.height:
                pic_area += sh.width * sh.height
        if getattr(sh, "has_chart", False) and sh.has_chart:
            counts["charts"] += 1
            try:
                chart_types.append(str(sh.chart.chart_type).split(".")[-1].split(" ")[0].lower())
            except Exception:
                chart_types.append("unknown")
        if getattr(sh, "has_table", False) and sh.has_table:
            counts["tables"] += 1
            table_dims.append(f"{len(sh.table.rows)}x{len(sh.table.columns)}")
        el = sh._element
        if el.tag.endswith("graphicFrame") and b"diagram" in (el.xml.encode() if hasattr(el, "xml") else b""):
            counts["smartart"] += 1
        if st == MSO_SHAPE_TYPE.LINE or el.tag.endswith("cxnSp"):
            counts["connectors"] += 1
        if st == MSO_SHAPE_TYPE.FREEFORM:
            counts["freeforms"] += 1
        if st == MSO_SHAPE_TYPE.AUTO_SHAPE:
            counts["autoshapes"] += 1
            try:
                name = str(sh.auto_shape_type).split(".")[-1].split(" ")[0].lower()
                ashape_types[name] += 1
                if any(k in name for k in ("arrow", "chevron", "pentagon")):
                    counts["arrows"] += 1
            except Exception:
                pass
            try:
                if sh.fill.type == 1 and sh.fill.fore_color and sh.fill.fore_color.type == 1:
                    colors[str(sh.fill.fore_color.rgb)] += 1
            except Exception:
                pass
        t = _text_of(sh)
        if t:
            counts["text"] += 1
            texts.append(t)
            text_boxes.append(sh)
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.name:
                        fonts[r.font.name] += 1
                    if r.font.size:
                        sizes.append(r.font.size.pt)
                    try:
                        if r.font.color and r.font.color.type == 1:
                            colors[str(r.font.color.rgb)] += 1
                    except Exception:
                        pass
    title = _title(slide)
    body = "\n".join(t for t in texts if t != title)
    notes = ""
    if slide.has_notes_slide:
        notes = (slide.notes_slide.notes_text_frame.text if slide.notes_slide.notes_text_frame else "").strip()
    alltext = f"{title}\n{body}"
    words = len(re.findall(r"\w+", alltext))
    feats = dict(
        slide_num=n, layout_name=slide.slide_layout.name, title=title[:500], body_text=body[:8000],
        notes=notes[:4000], word_count=words, n_shapes=counts["shapes"], n_text_shapes=counts["text"],
        n_pictures=counts["pictures"], picture_area_pct=round(100 * pic_area / (sw * shh), 1) if sw and shh else 0,
        n_charts=counts["charts"], chart_types=",".join(chart_types), n_tables=counts["tables"],
        table_dims=",".join(table_dims), n_groups=counts["groups"], n_smartart=counts["smartart"],
        n_connectors=counts["connectors"], n_autoshapes=counts["autoshapes"], n_arrow_shapes=counts["arrows"],
        n_freeforms=counts["freeforms"], autoshape_types=json.dumps(dict(ashape_types.most_common(8))),
        n_columns_est=_columns(text_boxes, sw), has_placeholder_text=int(bool(PLACEHOLDER_RE.search(alltext))),
        title_words=len(title.split()), min_font_pt=min(sizes) if sizes else None,
        max_font_pt=max(sizes) if sizes else None, fonts=json.dumps(dict(fonts.most_common(5))),
        colors=json.dumps(dict(colors.most_common(8))),
        margin_left_in=emu_to_in(left) if left is not None else None,
        margin_top_in=emu_to_in(top) if top is not None else None,
        margin_right_in=emu_to_in(sw - right) if right is not None else None,
        margin_bottom_in=emu_to_in(shh - bottom) if bottom is not None else None,
        hidden=int(slide._element.get("show") == "0"),
        quote_like=int(alltext.strip().startswith(("“", '"')) and words < 80),
    )
    return feats


def _layouts(prs, deck_id: str) -> list[tuple]:
    rows = []
    for mi, m in enumerate(prs.slide_masters):
        for li, lay in enumerate(m.slide_layouts):
            ph = []
            for p in lay.placeholders:
                try:
                    ph.append(dict(idx=p.placeholder_format.idx, type=str(p.placeholder_format.type).split(".")[-1].split(" ")[0],
                                   name=p.name, left_in=emu_to_in(p.left), top_in=emu_to_in(p.top),
                                   width_in=emu_to_in(p.width), height_in=emu_to_in(p.height)))
                except Exception:
                    continue
            rows.append((deck_id, mi, li, lay.name, json.dumps(ph)))
    return rows


# ---------- indexing ----------

def index_deck(db, cfg, source: str, root: Path, rel: Path, force: bool = False) -> str:
    path = root / rel
    did = deck_id_for(source, rel)
    md5 = file_md5(path)
    row = db.execute("SELECT md5, rendered FROM decks WHERE id=?", (did,)).fetchone()
    if row and row["md5"] == md5 and not force:
        return did
    keep_rendered = int(bool(row and row["md5"] == md5 and row["rendered"]))
    db.execute("DELETE FROM slides WHERE deck_id=?", (did,))
    db.execute("DELETE FROM slides_fts WHERE id LIKE ?", (did + "#%",))
    db.execute("DELETE FROM layouts WHERE deck_id=?", (did,))
    overrides = {r["slide_id"]: r for r in db.execute("SELECT * FROM overrides WHERE slide_id LIKE ?", (did + "#%",))}
    total, n_layouts, sw, shh = 0, 0, 0, 0
    for prs, numbered in _iter_deck(path):
        sw, shh = prs.slide_width, prs.slide_height
        _index_slides(db, did, source, numbered, sw, shh, overrides)
        total += len(numbered)
        if not db.execute("SELECT 1 FROM layouts WHERE deck_id=?", (did,)).fetchone():
            db.executemany("INSERT OR REPLACE INTO layouts VALUES (?,?,?,?,?)", _layouts(prs, did))
            n_layouts = sum(len(m.slide_layouts) for m in prs.slide_masters)
        db.commit()
    db.execute("INSERT OR REPLACE INTO decks (id,source,rel_path,name,md5,n_slides,width_in,height_in,n_layouts,indexed_at,rendered) "
               "VALUES (?,?,?,?,?,?,?,?,?,?,?)",
               (did, source, rel.as_posix(), rel.stem, md5, total, emu_to_in(sw), emu_to_in(shh),
                n_layouts, dt.datetime.now().isoformat(timespec="seconds"), keep_rendered))
    db.commit()
    print(f"indexed {did}: {total} slides")
    return did


BIG_DECK_BYTES = 150_000_000


def _iter_deck(path: Path):
    """Yield (presentation, [(slide_num, slide), ...]). Very large decks come in chunks."""
    import shutil
    import tempfile
    from .pptx_utils import chunk_plan, potx_to_pptx, slide_xml_bytes, subset_deck
    src = Path(path)
    tmp = Path(tempfile.mkdtemp(prefix="deckidx_"))
    if src.suffix.lower() == ".potx":
        src = potx_to_pptx(src, tmp / (src.stem + ".pptx"))
    try:
        if sum(slide_xml_bytes(src)) <= BIG_DECK_BYTES:
            prs = open_presentation(src)
            yield prs, list(enumerate(prs.slides, 1))
            return
        for chunk in chunk_plan(src):
            part = subset_deck(src, chunk, tmp / "chunk.pptx")
            prs = open_presentation(part)
            yield prs, list(zip(chunk, prs.slides))
            del prs
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def _index_slides(db, did, source, numbered, sw, shh, overrides):
    for n, slide in numbered:
        f = extract(slide, n, sw, shh)
        cat, sub, tags = classify(f)
        sid = f"{did}#{n}"
        if sid in overrides:
            o = overrides[sid]
            cat = o["category"] or cat
            sub = o["subcategory"] or sub
            if o["tags"]:
                tags = sorted(set(tags) | set(o["tags"].split(",")))
        f.pop("quote_like", None)
        f.update(id=sid, deck_id=did, source=source, category=cat, subcategory=sub, tags=",".join(tags),
                 assertion_title=int("assertion-title" in tags), curated_group=None,
                 thumb=f"thumbs/{did}/{n}.png")
        cols = ",".join(f.keys())
        db.execute(f"INSERT INTO slides ({cols}) VALUES ({','.join('?' * len(f))})", list(f.values()))
        db.execute("INSERT INTO slides_fts (id,title,body_text,notes,category,tags,layout_name) VALUES (?,?,?,?,?,?,?)",
                   (sid, f["title"], f["body_text"], f["notes"], cat, f["tags"], f["layout_name"]))


def _placeholder(path: Path) -> None:
    from PIL import Image, ImageDraw
    im = Image.new("RGB", (640, 360), "#eeeeee")
    ImageDraw.Draw(im).text((20, 170), "Render failed in LibreOffice. Open slide in PowerPoint.", fill="black")
    im.save(path)


def render_deck_thumbs(db, cfg, did: str, force: bool = False) -> None:
    from .render import contact_sheet, render_deck
    d = db.execute("SELECT * FROM decks WHERE id=?", (did,)).fetchone()
    if d["rendered"] and not force:
        return
    import shutil
    import tempfile
    from .pptx_utils import chunk_plan, potx_to_pptx, slide_xml_bytes, subset_deck
    root = C.source_roots(cfg)[d["source"]]
    out = C.library_root(cfg) / "thumbs" / did
    src = root / d["rel_path"]
    # Always render through subset_deck: it unhides hidden slides so page n is slide n,
    # and chunking bounds LibreOffice memory on very large decks. Existing thumbnails are
    # kept, so an interrupted run resumes. A chunk that crashes LibreOffice is split in half
    # until the failing slide is isolated; that slide gets a placeholder thumbnail.
    out.mkdir(parents=True, exist_ok=True)
    tmp = Path(tempfile.mkdtemp(prefix="deckrender_"))
    if src.suffix.lower() == ".potx":
        src = potx_to_pptx(src, tmp / (src.stem + ".pptx"))
    failed: list[int] = []

    def go(chunk: list[int]) -> None:
        chunk = [c for c in chunk if force or not (out / f"{c}.png").exists()]
        if not chunk:
            return
        part = subset_deck(src, chunk, tmp / "part.pptx")
        try:
            for p in render_deck(part, tmp / "r", cfg):
                shutil.move(str(p), out / f"{chunk[int(p.stem) - 1]}.png")
        except Exception:
            if len(chunk) == 1:
                failed.append(chunk[0])
                _placeholder(out / f"{chunk[0]}.png")
                return
            mid = len(chunk) // 2
            go(chunk[:mid])
            go(chunk[mid:])

    budget = 40_000_000 if sum(slide_xml_bytes(src)) > BIG_DECK_BYTES else 20_000_000
    for ci, chunk in enumerate(chunk_plan(src, budget)):
        go(chunk)
        print(f"  slides {chunk[0]}-{chunk[-1]} done")
    shutil.rmtree(tmp, ignore_errors=True)
    if failed:
        print(f"  render failed for slides {failed}; placeholder thumbnails written")
    pngs = sorted(out.glob("*.png"), key=lambda p: int(p.stem))
    sheets = C.library_root(cfg) / "sheets" / did
    for k in range(0, len(pngs), 48):
        contact_sheet(pngs[k:k + 48], sheets / f"{k // 48 + 1}.jpg")
    db.execute("UPDATE decks SET rendered=1 WHERE id=?", (did,))
    db.commit()
    print(f"rendered {did}: {len(pngs)} thumbnails")


def import_curated(db, cfg) -> int:
    """Mark slides hand-picked in TEMPLATE_LIBRARY_ROOT/claude_design_inputs/selection_index.csv."""
    idx = Path(cfg["TEMPLATE_LIBRARY_ROOT"]) / "claude_design_inputs" / "selection_index.csv"
    if not idx.exists():
        return 0
    n = 0
    with open(idx, newline="") as fh:
        for r in csv.DictReader(fh):
            group = Path(r["output_pdf"]).stem
            stem = Path(r["source_pdf"]).stem
            deck = db.execute("SELECT id FROM decks WHERE source='slideworks' AND name=?", (stem,)).fetchone()
            if not deck:  # chunked PDFs such as due-diligence_00.pdf
                key = re.sub(r"_\d+$", "", stem).replace("-", "_")
                deck = db.execute("SELECT id FROM decks WHERE source='slideworks' AND lower(name) LIKE ?",
                                  (f"%{key}%",)).fetchone()
            if not deck:
                continue
            num = int(r["orig_slide_num"])
            want = (r.get("title") or "").strip().lower()[:30]
            # verify by title; slide numbers can drift between template versions
            cand = db.execute("SELECT id, slide_num, lower(title) t FROM slides WHERE deck_id=? AND slide_num BETWEEN ? AND ?",
                              (deck["id"], num - 12, num + 12)).fetchall()
            hit = next((c for c in cand if c["slide_num"] == num and want and c["t"].startswith(want[:20])), None)
            hit = hit or next((c for c in sorted(cand, key=lambda c: abs(c["slide_num"] - num)) if want and c["t"].startswith(want[:20])), None)
            if not hit and want:
                allc = db.execute("SELECT id, slide_num, lower(title) t FROM slides WHERE deck_id=?", (deck["id"],)).fetchall()
                hit = next((c for c in sorted(allc, key=lambda c: abs(c["slide_num"] - num)) if c["t"].startswith(want[:20])), None)
            hit = hit or next((c for c in cand if c["slide_num"] == num), None)
            if hit:
                db.execute("UPDATE slides SET curated_group=? WHERE id=?", (group, hit["id"]))
                n += 1
    db.commit()
    return n


def export_csv(db, cfg) -> Path:
    out = C.library_root(cfg) / "catalog.csv"
    rows = db.execute("SELECT * FROM slides ORDER BY deck_id, slide_num").fetchall()
    with open(out, "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(rows[0].keys())
        for r in rows:
            w.writerow(list(r))
    return out


def build(cfg, render: bool = True, force: bool = False, only: str | None = None) -> None:
    db = connect(cfg)
    seen = set()
    for source, root, rel in find_decks(cfg):
        did = deck_id_for(source, rel)
        if only and only not in did:
            continue
        seen.add(did)
        try:
            index_deck(db, cfg, source, root, rel, force)
            if render:
                render_deck_thumbs(db, cfg, did, force)
        except Exception as e:  # keep going; report at the end
            print(f"ERROR {did}: {e}")
    if not only:
        stale = [r["id"] for r in db.execute("SELECT id FROM decks") if r["id"] not in seen]
        for did in stale:
            db.execute("DELETE FROM slides WHERE deck_id=?", (did,))
            db.execute("DELETE FROM layouts WHERE deck_id=?", (did,))
            db.execute("DELETE FROM decks WHERE id=?", (did,))
            print(f"removed stale deck {did}")
    print(f"curated slides matched: {import_curated(db, cfg)}")
    db.commit()
    print(f"catalog: {export_csv(db, cfg)}")


def stats(cfg) -> str:
    db = connect(cfg)
    lines = []
    for r in db.execute("SELECT source, count(*) decks, sum(n_slides) slides, sum(rendered) rendered FROM decks GROUP BY source"):
        lines.append(f"{r['source']}: {r['decks']} decks, {r['slides']} slides, {r['rendered']} rendered")
    lines.append("")
    for r in db.execute("SELECT category, count(*) n FROM slides WHERE hidden=0 GROUP BY category ORDER BY n DESC"):
        lines.append(f"{r['category']:<20} {r['n']}")
    return "\n".join(lines)


def search(cfg, query: str | None = None, category: str | None = None, source: str | None = None,
           chart: str | None = None, curated: bool = False, limit: int = 20) -> list[dict]:
    db = connect(cfg)
    where, args = ["s.hidden=0", "s.category NOT IN ('instructions')"], []
    join = ""
    if query:
        join = "JOIN slides_fts f ON f.id = s.id"
        where.append("slides_fts MATCH ?")
        args.append(query)
    if category:
        where.append("s.category = ?")
        args.append(category)
    if source:
        where.append("s.source = ?")
        args.append(source)
    if chart:
        where.append("s.chart_types LIKE ?")
        args.append(f"%{chart}%")
    if curated:
        where.append("s.curated_group IS NOT NULL")
    order = "ORDER BY (s.curated_group IS NOT NULL) DESC, s.assertion_title DESC, s.word_count"
    sql = f"SELECT s.* FROM slides s {join} WHERE {' AND '.join(where)} {order} LIMIT ?"
    rows = db.execute(sql, args + [limit]).fetchall()
    root = C.library_root(cfg)
    srcs = C.source_roots(cfg)
    out = []
    for r in rows:
        d = db.execute("SELECT rel_path FROM decks WHERE id=?", (r["deck_id"],)).fetchone()
        out.append(dict(id=r["id"], category=r["category"], sub=r["subcategory"], title=r["title"][:90],
                        words=r["word_count"], deck=str(srcs[r["source"]] / d["rel_path"]),
                        slide=r["slide_num"], thumb=str(root / r["thumb"]), curated=r["curated_group"]))
    return out


def tag(cfg, slide_id: str, category: str | None, tags: str | None, note: str | None, quality: int | None) -> None:
    db = connect(cfg)
    db.execute("INSERT INTO overrides (slide_id,category,tags,note,quality) VALUES (?,?,?,?,?) "
               "ON CONFLICT(slide_id) DO UPDATE SET category=coalesce(excluded.category,category), "
               "tags=coalesce(excluded.tags,tags), note=coalesce(excluded.note,note), quality=coalesce(excluded.quality,quality)",
               (slide_id, category, tags, note, quality))
    if category:
        db.execute("UPDATE slides SET category=? WHERE id=?", (category, slide_id))
    db.commit()
